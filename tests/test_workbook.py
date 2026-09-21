"""Tests for workbook.py (OutputWorkbook)."""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from kosaccounts.models import BankRow, PurchaseRow
from kosaccounts.workbook import OutputWorkbook


def make_row(
    *,
    day: int = 1,
    company: str = "Mood Fabrics",
    net: str = "100.00",
    sales_tax: str = "8.88",
    total: str = "108.88",
    status: str = "OK",
    notes: str = "",
    bank_ref: str = "",
) -> PurchaseRow:
    return PurchaseRow(
        date=date(2026, 9, day),
        company=company,
        category="Fabric",
        net=Decimal(net),
        sales_tax=Decimal(sales_tax),
        total=Decimal(total),
        tax_rate=Decimal("8.875"),
        payment_method="VISA ****1234",
        notes=notes,
        year=2026,
        source_file=f"expenses/2026-09-0{day} - Mood.pdf",
        processed_on=datetime(2026, 9, day, 10, 30, 0),
        status=status,
        bank_ref=bank_ref,
    )


def make_bank_row(
    *,
    day: int = 1,
    description: str = "MOOD FABRICS NYC",
    amount: str = "108.88",
    direction: str = "debit",
    status: str = "No receipt",
    matched_source: str = "",
    ref: str = "ref-1",
) -> BankRow:
    return BankRow(
        date=date(2026, 9, day),
        description=description,
        amount=Decimal(amount),
        direction=direction,
        status=status,
        matched_source=matched_source,
        ref=ref,
        statement_file="bank/statement_sep.html",
        statement_start=date(2026, 9, 1),
        statement_end=date(2026, 9, 30),
        processed_on=datetime(2026, 9, day, 10, 30, 0),
    )


class TestOutputWorkbook:
    def test_create_append_save_reopen_roundtrip(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)

        row1 = make_row(day=1, company="Mood Fabrics")
        row2 = make_row(day=2, company="B&H Photo")
        wb.append([row1, row2])
        assert row1.sheet_row == 4
        assert row2.sheet_row == 5
        wb.save()

        assert path.exists()
        assert not path.with_suffix(".xlsx.tmp").exists()

        wb2 = OutputWorkbook(path)
        rows = wb2.rows()
        assert len(rows) == 2

        for original, reread in zip([row1, row2], rows):
            assert reread.date == original.date
            assert reread.company == original.company
            assert reread.category == original.category
            assert reread.net == round(Decimal(original.net), 2)
            assert reread.sales_tax == round(Decimal(original.sales_tax), 2)
            assert reread.total == round(Decimal(original.total), 2)
            assert reread.payment_method == original.payment_method
            assert reread.notes == original.notes
            assert reread.year == original.year
            assert reread.source_file == original.source_file
            assert reread.processed_on == original.processed_on
            assert reread.status == original.status
            assert reread.bank_ref == original.bank_ref
            assert reread.sheet_row is not None

    def test_append_after_reopen_and_totals_row(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        wb.append([make_row(day=1), make_row(day=2)])
        wb.save()

        wb2 = OutputWorkbook(path)
        wb2.append([make_row(day=3, company="Amazon.com")])
        wb2.save()

        rows = wb2.rows()
        assert len(rows) == 3
        assert [r.sheet_row for r in rows] == [4, 5, 6]

        raw = openpyxl.load_workbook(path, data_only=False)
        ws = raw[PurchaseRow.SHEET_NAME]

        # totals row sits directly below the last data row
        assert ws.cell(row=7, column=2).value == "Total"
        assert ws.cell(row=7, column=5).value == "=SUBTOTAL(109,E4:E6)"
        assert ws.cell(row=7, column=6).value == "=SUBTOTAL(109,F4:F6)"
        assert ws.cell(row=7, column=7).value == "=SUBTOTAL(109,G4:G6)"

        table = ws.tables[PurchaseRow.TABLE_NAME]
        assert table.ref == "B3:P6"

        # Check formula text present in J4 (Net=E, Sales Tax=F, Total=G, Check=J)
        assert ws.cell(row=4, column=10).value == '=IF(ROUND(E4+F4,2)=ROUND(G4,2),"ok","!!!!")'

    def test_update_changes_fields_and_review_fill(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        row1 = make_row(day=1)
        row2 = make_row(day=2)
        wb.append([row1, row2])
        wb.save()

        row1.bank_ref = "TXN-123"
        row1.notes = "Matched to bank statement"
        row2.status = "Review"
        row2.notes = "No bank transaction found"
        wb.update([row1, row2])
        wb.save()

        wb2 = OutputWorkbook(path)
        rows = wb2.rows()
        assert rows[0].bank_ref == "TXN-123"
        assert rows[0].notes == "Matched to bank statement"
        assert rows[1].status == "Review"
        assert rows[1].notes == "No bank transaction found"

        raw = openpyxl.load_workbook(path, data_only=False)
        ws = raw[PurchaseRow.SHEET_NAME]
        review_fill = ws.cell(row=5, column=2).fill
        assert review_fill.fgColor.rgb in ("00FFF2CC", "FFFFF2CC", "FFF2CC")
        ok_fill = ws.cell(row=4, column=2).fill
        assert ok_fill.fill_type is None

    def test_superseded_row_gets_grey_fill_and_strike_font(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        row1 = make_row(day=1)
        row2 = make_row(day=2)
        wb.append([row1, row2])
        wb.save()

        row1.status = "Superseded"
        row1.notes = "Superseded by re-upload on 2026-09-18"
        wb.update([row1])
        wb.save()

        raw = openpyxl.load_workbook(path, data_only=False)
        ws = raw[PurchaseRow.SHEET_NAME]

        superseded_fill = ws.cell(row=4, column=2).fill
        assert superseded_fill.fgColor.rgb in ("00EDEDED", "FFEDEDED", "EDEDED")
        superseded_font = ws.cell(row=4, column=2).font
        assert superseded_font.strike is True
        assert superseded_font.color.rgb in ("00808080", "FF808080", "808080")

        # The untouched OK row keeps no fill and a normal (non-strike) font.
        ok_fill = ws.cell(row=5, column=2).fill
        assert ok_fill.fill_type is None
        ok_font = ws.cell(row=5, column=2).font
        assert not ok_font.strike

    def test_update_without_sheet_row_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        row = make_row()
        with pytest.raises(ValueError):
            wb.update([row])

    def test_save_leaves_no_tmp_file(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        wb.append([make_row()])
        wb.save()
        tmp = path.with_suffix(".xlsx.tmp")
        assert not tmp.exists()

    def test_zero_row_workbook_saves_and_reopens_empty(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        wb.save()

        assert path.exists()
        assert not path.with_suffix(".xlsx.tmp").exists()

        wb2 = OutputWorkbook(path)
        assert wb2.rows() == []

    def test_table_style_is_unstriped(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        wb.append([make_row()])
        wb.save()

        raw = openpyxl.load_workbook(path, data_only=False)
        ws = raw[PurchaseRow.SHEET_NAME]
        table = ws.tables[PurchaseRow.TABLE_NAME]
        assert table.tableStyleInfo.name == "TableStyleLight1"
        assert table.tableStyleInfo.showRowStripes is False
        assert table.tableStyleInfo.showColumnStripes is False

        header_font = ws.cell(row=3, column=2).font
        assert header_font.bold is True


def test_save_creates_missing_output_directory(tmp_path):
    """A fresh install has no output/ folder yet; save() must create it."""
    from kosaccounts.workbook import OutputWorkbook

    target = tmp_path / "output" / "nested" / "kosibah_import.xlsx"
    wb = OutputWorkbook(target)
    wb.save()
    assert target.exists()
    assert OutputWorkbook(target).rows() == []


class TestBankSheet:
    def test_bank_sheet_not_created_until_needed(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        wb.append([make_row()])
        wb.save()

        assert wb.bank_rows() == []

        raw = openpyxl.load_workbook(path, data_only=False)
        assert BankRow.SHEET_NAME not in raw.sheetnames

    def test_append_bank_creates_sheet_after_purchases(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        wb.append([make_row()])

        br1 = make_bank_row(day=1, ref="ref-1")
        br2 = make_bank_row(day=2, ref="ref-2", status="Matched", matched_source="expenses/x.pdf")
        wb.append_bank([br1, br2])
        assert br1.sheet_row == 4
        assert br2.sheet_row == 5
        wb.save()

        raw = openpyxl.load_workbook(path, data_only=False)
        assert raw.sheetnames.index(BankRow.SHEET_NAME) > raw.sheetnames.index(PurchaseRow.SHEET_NAME)

        ws = raw[BankRow.SHEET_NAME]
        assert ws.cell(row=1, column=2).value == "Kosibah LLC bank transactions - reconciliation"
        for i, header in enumerate(BankRow.COLUMNS):
            assert ws.cell(row=3, column=2 + i).value == header
        assert ws.freeze_panes == "B4"

        table = ws.tables[BankRow.TABLE_NAME]
        assert table.ref == "B3:L5"
        assert table.tableStyleInfo.name == "TableStyleLight1"
        assert table.tableStyleInfo.showRowStripes is False

    def test_bank_rows_roundtrip(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        row = make_bank_row(day=5, description="AMZ MKTP", amount="42.00", ref="ref-9")
        wb.append_bank([row])
        wb.save()

        wb2 = OutputWorkbook(path)
        rows = wb2.bank_rows()
        assert len(rows) == 1
        reread = rows[0]
        assert reread.date == row.date
        assert reread.description == row.description
        assert reread.amount == round(Decimal(row.amount), 2)
        assert reread.direction == row.direction
        assert reread.status == row.status
        assert reread.matched_source == row.matched_source
        assert reread.ref == row.ref
        assert reread.statement_file == row.statement_file
        assert reread.statement_start == row.statement_start
        assert reread.statement_end == row.statement_end
        assert reread.processed_on == row.processed_on
        assert reread.sheet_row == 4

    def test_update_bank_by_sheet_row(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        row = make_bank_row(ref="ref-1", status="No receipt")
        wb.append_bank([row])
        wb.save()

        row.status = "Matched"
        row.matched_source = "expenses/2026-09-01 - Mood.pdf"
        wb.update_bank([row])
        wb.save()

        wb2 = OutputWorkbook(path)
        reread = wb2.bank_rows()[0]
        assert reread.status == "Matched"
        assert reread.matched_source == "expenses/2026-09-01 - Mood.pdf"

    def test_update_bank_without_sheet_row_raises(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        row = make_bank_row()
        with pytest.raises(ValueError):
            wb.update_bank([row])

    def test_bank_status_fills(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        no_receipt = make_bank_row(day=1, ref="a", status="No receipt")
        ignored = make_bank_row(day=2, ref="b", status="Ignored")
        matched = make_bank_row(day=3, ref="c", status="Matched", matched_source="expenses/x.pdf")
        wb.append_bank([no_receipt, ignored, matched])
        wb.save()

        raw = openpyxl.load_workbook(path, data_only=False)
        ws = raw[BankRow.SHEET_NAME]

        no_receipt_fill = ws.cell(row=4, column=2).fill
        assert no_receipt_fill.fgColor.rgb in ("00FFF2CC", "FFFFF2CC", "FFF2CC")

        ignored_fill = ws.cell(row=5, column=2).fill
        assert ignored_fill.fgColor.rgb in ("00EDEDED", "FFEDEDED", "EDEDED")

        matched_fill = ws.cell(row=6, column=2).fill
        assert matched_fill.fill_type is None

    def test_bank_sheet_survives_reopen_alongside_purchases(self, tmp_path: Path) -> None:
        path = tmp_path / "kosibah_import.xlsx"
        wb = OutputWorkbook(path)
        wb.append([make_row()])
        wb.append_bank([make_bank_row()])
        wb.save()

        wb2 = OutputWorkbook(path)
        assert len(wb2.rows()) == 1
        assert len(wb2.bank_rows()) == 1

        # Appending more of each keeps both sheets independent.
        wb2.append([make_row(day=2)])
        wb2.append_bank([make_bank_row(day=2, ref="ref-2")])
        wb2.save()

        wb3 = OutputWorkbook(path)
        assert len(wb3.rows()) == 2
        assert len(wb3.bank_rows()) == 2
