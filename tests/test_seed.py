"""Tests for kosaccounts.seed.seed_from_master."""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from kosaccounts.seed import SeedStats, seed_from_master

REPO_ROOT = Path(__file__).resolve().parent.parent
REAL_MASTER = REPO_ROOT / "specs" / "Kosibah LLC purchases and receipts.xlsx"

# (Date, Company, Category, Schedule C, Net, Sales Tax, Total, Tax rate, Payment method, Check, Notes, Year)
_PURCHASE_ROWS = [
    (datetime(2024, 1, 1), "Pacific Trimming", "Haberdashery", "Cost of Goods Sold", 10.0, 0, 10.0, 0, "Cash", "ok", None, 2024),
    (datetime(2024, 1, 5), "T Mobile", "Telephone", "Other expenses", 20.0, 0, 20.0, 0, "Cheque 0104", "ok", None, 2024),
    (datetime(2024, 2, 1), "T-Mobile", "Telephone", "Other expenses", 25.0, 0, 25.0, 0, "Mastercard", "ok", None, 2024),
    (datetime(2024, 2, 5), "T-Mobile", "Telephone", "Other expenses", 30.0, 0, 30.0, 0, "Mastercard", "ok", None, 2024),
    (datetime(2024, 3, 1), "Amazon.com", "Office supplies", "Office expense", 15.0, 0, 15.0, 0, "Discover", "ok", None, 2024),
    (datetime(2024, 3, 5), "Amazon.com", "Office equipment", "Office expense", 45.0, 0, 45.0, 0, "Discover", "ok", None, 2024),
    (datetime(2024, 4, 1), "Guide Fabrics", "Fabric", "Cost of Goods Sold", 100.0, 0, 100.0, 0, "Wire transfer", "ok", None, 2024),
    (datetime(2024, 1, 10), "Pacific Trimming", "Haberdashery", "Cost of Goods Sold", 5.0, 0, 5.0, 0, "Cheque", "ok", None, 2024),
]

# (Original, Schedule C) -- deliberately not in the order categories are used in _PURCHASE_ROWS,
# and including one category ("Bank charges") never used by a purchase row.
_LOOKUP_ROWS = [
    ("Fabric", "Cost of Goods Sold"),
    ("Haberdashery", "Cost of Goods Sold"),
    ("Office equipment", "Office expense"),
    ("Office supplies", "Office expense"),
    ("Telephone", "Other expenses"),
    ("Bank charges", "Other expenses"),
]


def _build_master_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchases"

    ws["B1"] = "Kosibah Purchases (test)"

    headers = [
        "Date", "Company", "Category", "Schedule C", "Net", "Sales Tax",
        "Total", "Tax rate", "Payment method", "Check", "Notes", "Year",
    ]
    for offset, header in enumerate(headers):
        ws.cell(row=3, column=2 + offset, value=header)

    row_num = 4
    net_total = 0.0
    tax_total = 0.0
    total_total = 0.0
    for values in _PURCHASE_ROWS:
        for offset, value in enumerate(values):
            ws.cell(row=row_num, column=2 + offset, value=value)
        net_total += values[4]
        tax_total += values[5]
        total_total += values[6]
        row_num += 1

    # Total row: "Total" in column B, sums in Net/Sales Tax/Total.
    ws.cell(row=row_num, column=2, value="Total")
    ws.cell(row=row_num, column=6, value=net_total)
    ws.cell(row=row_num, column=7, value=tax_total)
    ws.cell(row=row_num, column=8, value=total_total)
    ws.cell(row=row_num, column=11, value="ok")

    lookup = wb.create_sheet("Lookup values")
    lookup["F2"] = "Original"
    lookup["G2"] = "Schedule C"
    for offset, (original, schedule_c) in enumerate(_LOOKUP_ROWS):
        lookup.cell(row=3 + offset, column=6, value=original)
        lookup.cell(row=3 + offset, column=7, value=schedule_c)

    wb.save(path)


def _read_csv(path: Path) -> list[dict]:
    with path.open("r", newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


@pytest.fixture
def master_path(tmp_path: Path) -> Path:
    path = tmp_path / "master.xlsx"
    _build_master_workbook(path)
    return path


@pytest.fixture
def data_dir(tmp_path: Path) -> Path:
    d = tmp_path / "data"
    d.mkdir()
    return d


def test_seed_stats(master_path: Path, data_dir: Path) -> None:
    stats = seed_from_master(master_path, data_dir)
    assert isinstance(stats, SeedStats)
    assert stats.purchase_rows == 8
    # 5 raw supplier names, 2 fold together (T Mobile / T-Mobile) -> 4 canonical suppliers.
    assert stats.suppliers == 4
    assert stats.categories == 6
    # Cash, Cheque (merged), Mastercard, Discover, Wire transfer
    assert stats.payment_methods == 5
    assert stats.aliases_added == 1


def test_categories_csv_matches_lookup_table_in_file_order(master_path: Path, data_dir: Path) -> None:
    seed_from_master(master_path, data_dir)
    rows = _read_csv(data_dir / "categories.csv")
    assert [r["Category"] for r in rows] == [orig for orig, _ in _LOOKUP_ROWS]
    by_name = {r["Category"]: r["Schedule C"] for r in rows}
    assert by_name == dict(_LOOKUP_ROWS)


def test_suppliers_csv_default_category_and_counts(master_path: Path, data_dir: Path) -> None:
    seed_from_master(master_path, data_dir)
    rows = {r["Supplier"]: r for r in _read_csv(data_dir / "suppliers.csv")}

    # T Mobile / T-Mobile collide; T-Mobile (2 rows) is more frequent than T Mobile (1 row),
    # so T-Mobile is canonical and must be the only row in suppliers.csv.
    assert "T-Mobile" in rows
    assert "T Mobile" not in rows
    assert rows["T-Mobile"]["Default category"] == "Telephone"
    assert rows["T-Mobile"]["Count"] == "3"
    assert rows["T-Mobile"]["Last date"] == "2024-02-05"

    # Pacific Trimming: two rows, both "Haberdashery" -> unambiguous default, count 2.
    assert rows["Pacific Trimming"]["Default category"] == "Haberdashery"
    assert rows["Pacific Trimming"]["Count"] == "2"
    assert rows["Pacific Trimming"]["Last date"] == "2024-01-10"

    # Amazon.com: one "Office supplies" row (2024-03-01) and one "Office equipment" row
    # (2024-03-05, more recent) -> tied 1/1, tie-break picks the most recent row's category.
    assert rows["Amazon.com"]["Default category"] == "Office equipment"
    assert rows["Amazon.com"]["Other categories"] == "Office supplies"
    assert rows["Amazon.com"]["Count"] == "2"

    assert rows["Guide Fabrics"]["Default category"] == "Fabric"
    assert rows["Guide Fabrics"]["Count"] == "1"


def test_payment_methods_csv_strips_cheque_numbers_and_merges(master_path: Path, data_dir: Path) -> None:
    seed_from_master(master_path, data_dir)
    rows = {r["Payment method"]: int(r["Count"]) for r in _read_csv(data_dir / "payment_methods.csv")}
    assert "Cheque 0104" not in rows
    assert rows["Cheque"] == 2  # "Cheque 0104" + plain "Cheque"
    assert rows["Cash"] == 1
    assert rows["Mastercard"] == 2
    assert rows["Discover"] == 2
    assert rows["Wire transfer"] == 1


def test_supplier_aliases_csv_seeded_with_collision(master_path: Path, data_dir: Path) -> None:
    seed_from_master(master_path, data_dir)
    rows = _read_csv(data_dir / "supplier_aliases.csv")
    assert {"Alias": "T Mobile", "Supplier": "T-Mobile"} in rows


def test_supplier_aliases_merge_preserves_hand_added_rows(master_path: Path, data_dir: Path) -> None:
    # First seed run.
    seed_from_master(master_path, data_dir)

    # Simulate a hand-added alias row that the pipeline would never detect on its own.
    aliases_path = data_dir / "supplier_aliases.csv"
    rows = _read_csv(aliases_path)
    rows.append({"Alias": "Amazon", "Supplier": "Amazon.com"})
    with aliases_path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=["Alias", "Supplier"])
        writer.writeheader()
        writer.writerows(rows)

    # Re-seed: the hand-added row must survive, and no duplicate detected alias is added.
    stats = seed_from_master(master_path, data_dir)
    assert stats.aliases_added == 0

    rows_after = _read_csv(aliases_path)
    assert {"Alias": "Amazon", "Supplier": "Amazon.com"} in rows_after
    assert {"Alias": "T Mobile", "Supplier": "T-Mobile"} in rows_after
    assert len(rows_after) == 2


def test_suppliers_csv_written_atomically_no_leftover_tmp_file(master_path: Path, data_dir: Path) -> None:
    seed_from_master(master_path, data_dir)
    assert (data_dir / "suppliers.csv").exists()
    assert not any(data_dir.glob("*.tmp"))


@pytest.mark.skipif(not REAL_MASTER.exists(), reason="real master workbook not available")
def test_seed_from_real_master(tmp_path: Path) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    stats = seed_from_master(REAL_MASTER, data_dir)
    assert stats.purchase_rows == 1438
    # NOTE: the plan/task brief states 236 suppliers, which matches the 236 distinct raw
    # "Company" strings in the Purchases sheet (verified directly). But the contract also
    # requires folding spelling collisions ("T Mobile"/"T-Mobile", "Wordpress.com"/
    # "WordPress.com") into their canonical spelling and NOT giving the alias its own row in
    # suppliers.csv. With those 2 known collisions folded, suppliers.csv ends up with 234
    # canonical rows. Asserting the post-fold count that suppliers.csv (and this SeedStats)
    # actually produces; see the final report for this as a flagged contract discrepancy.
    assert stats.suppliers == 234
    # NOTE: the plan/task brief states 30 categories, but the "Lookup values" sheet's
    # Original/Schedule C table (F2:G31, header row 2) contains 29 data rows (F3:F31) in the
    # real workbook as of 2026-09-18, and all 29 are the only distinct categories used across
    # the 1438 purchase rows. Asserting the verified actual count rather than the stated 30;
    # see the final report for this as a flagged contract/spec discrepancy.
    assert stats.categories == 29
