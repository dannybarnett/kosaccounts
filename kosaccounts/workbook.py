"""Output workbook (openpyxl): output/kosibah_import.xlsx.

Sheet "Purchases" (table "KosibahImport"): row 1 title ("Kosibah LLC purchases - imported"), row 3
headers (PurchaseRow.COLUMNS, columns B..P), data from row 4, a totals row after the last data row
(Net, Sales Tax, Total summed with SUBTOTAL(109, ...) so filters work). The Excel table spans
header..last data row (totals row is outside the table, like the master's).

Sheet "Bank" (table "KosibahBank"), created the first time it is needed: row 1 title ("Kosibah LLC
bank transactions - reconciliation"), row 3 headers (BankRow.COLUMNS, columns B..L), data from row
4. No totals row.

Formats: Date yyyy-mm-dd; Net/Sales Tax/Total #,##0.00; Tax rate 0.000; Processed on yyyy-mm-dd
hh:mm. Check column formula per row: =IF(ROUND(E{r}+F{r},2)=ROUND(G{r},2),"ok","!!!!").
Status == "Review" -> whole row filled amber (FFF2CC). Status == "Superseded" -> whole row filled
grey (EDEDED) with strike-through font. Status == "OK" -> no fill, normal font.
Save is atomic: write to <path>.tmp then os.replace.

CONTRACT (implement; do not change signatures):
"""

from __future__ import annotations

import os
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from kosaccounts.models import BankRow, PurchaseRow

TITLE = "Kosibah LLC purchases - imported"
HEADER_ROW = 3
FIRST_DATA_ROW = 4
FIRST_COL = 2  # column B
LAST_COL = 16  # column P

BANK_TITLE = "Kosibah LLC bank transactions - reconciliation"
BANK_HEADER_ROW = 3
BANK_FIRST_DATA_ROW = 4
BANK_FIRST_COL = 2  # column B
BANK_LAST_COL = 12  # column L

REVIEW_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
SUPERSEDED_FILL = PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")
NO_FILL = PatternFill(fill_type=None)
SUPERSEDED_FONT = Font(strike=True, color="808080")
NORMAL_FONT = Font()

_TABLE_STYLE = TableStyleInfo(
    name="TableStyleLight1",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=False,
    showColumnStripes=False,
)

COLUMN_WIDTHS = {
    "B": 12,
    "C": 32,
    "D": 22,
    "E": 12,
    "F": 12,
    "G": 12,
    "H": 12,
    "I": 20,
    "J": 8,
    "K": 40,
    "L": 6,
    "M": 44,
    "N": 18,
    "O": 10,
    "P": 24,
}

def _check_formula(r: int) -> str:
    return f'=IF(ROUND(E{r}+F{r},2)=ROUND(G{r},2),"ok","!!!!")'


class OutputWorkbook:
    def __init__(self, path: Path) -> None:
        """Open if exists, else create in memory (file written on save())."""
        self.path = Path(path)
        if self.path.exists():
            self.wb = openpyxl.load_workbook(self.path, data_only=False)
            self.ws = self.wb[PurchaseRow.SHEET_NAME]
            self._last_data_row = self._find_last_data_row()
            if BankRow.SHEET_NAME in self.wb.sheetnames:
                self.ws_bank = self.wb[BankRow.SHEET_NAME]
                self._last_bank_row = self._find_last_bank_row()
            else:
                self.ws_bank = None
                self._last_bank_row = None
        else:
            self.wb = openpyxl.Workbook()
            default_sheet = self.wb.active
            self.ws = self.wb.create_sheet(PurchaseRow.SHEET_NAME)
            if default_sheet.title != self.ws.title:
                self.wb.remove(default_sheet)
            self.ws.cell(row=1, column=FIRST_COL, value=TITLE)
            for i, header in enumerate(PurchaseRow.COLUMNS):
                self.ws.cell(row=HEADER_ROW, column=FIRST_COL + i, value=header)
            self._last_data_row = HEADER_ROW  # no data rows yet
            self.ws_bank = None
            self._last_bank_row = None

    # ------------------------------------------------------------------
    # Purchases sheet
    # ------------------------------------------------------------------

    def _find_last_data_row(self) -> int:
        """Row 4 down to the row before "Total" in column B, or the last non-empty B."""
        row = FIRST_DATA_ROW
        last = HEADER_ROW  # no data rows sentinel
        while True:
            val = self.ws.cell(row=row, column=FIRST_COL).value
            if val is None:
                break
            if isinstance(val, str) and val.strip() == "Total":
                break
            last = row
            row += 1
        return last

    def _has_data(self) -> bool:
        return self._last_data_row >= FIRST_DATA_ROW

    def rows(self) -> list[PurchaseRow]:
        """All data rows, with sheet_row set."""
        result: list[PurchaseRow] = []
        if not self._has_data():
            return result
        for r in range(FIRST_DATA_ROW, self._last_data_row + 1):
            row_date = self.ws.cell(row=r, column=2).value
            if isinstance(row_date, datetime):
                row_date = row_date.date()
            company = self.ws.cell(row=r, column=3).value or ""
            category = self.ws.cell(row=r, column=4).value or ""
            net = self.ws.cell(row=r, column=5).value
            sales_tax = self.ws.cell(row=r, column=6).value
            total = self.ws.cell(row=r, column=7).value
            tax_rate = self.ws.cell(row=r, column=8).value
            payment_method = self.ws.cell(row=r, column=9).value or ""
            notes = self.ws.cell(row=r, column=11).value or ""
            year = self.ws.cell(row=r, column=12).value
            source_file = self.ws.cell(row=r, column=13).value or ""
            processed_on = self.ws.cell(row=r, column=14).value
            status = self.ws.cell(row=r, column=15).value or "OK"
            bank_ref = self.ws.cell(row=r, column=16).value
            if bank_ref is None:
                bank_ref = ""

            row = PurchaseRow(
                date=row_date,
                company=str(company),
                category=str(category),
                net=Decimal(str(net)) if net is not None else Decimal("0"),
                sales_tax=Decimal(str(sales_tax)) if sales_tax is not None else Decimal("0"),
                total=Decimal(str(total)) if total is not None else Decimal("0"),
                tax_rate=Decimal(str(tax_rate)) if tax_rate is not None else Decimal("0"),
                payment_method=str(payment_method),
                notes=str(notes),
                year=int(year) if year is not None else 0,
                source_file=str(source_file),
                processed_on=processed_on,
                status=status,
                bank_ref=str(bank_ref),
                sheet_row=r,
            )
            result.append(row)
        return result

    def _write_row(self, r: int, row: PurchaseRow) -> None:
        self.ws.cell(row=r, column=2, value=row.date)
        self.ws.cell(row=r, column=3, value=row.company)
        self.ws.cell(row=r, column=4, value=row.category)
        self.ws.cell(row=r, column=5, value=float(round(Decimal(row.net), 2)))
        self.ws.cell(row=r, column=6, value=float(round(Decimal(row.sales_tax), 2)))
        self.ws.cell(row=r, column=7, value=float(round(Decimal(row.total), 2)))
        self.ws.cell(row=r, column=8, value=float(row.tax_rate))
        self.ws.cell(row=r, column=9, value=row.payment_method)
        self.ws.cell(row=r, column=10, value=_check_formula(r))
        self.ws.cell(row=r, column=11, value=row.notes)
        self.ws.cell(row=r, column=12, value=row.year)
        self.ws.cell(row=r, column=13, value=row.source_file)
        self.ws.cell(row=r, column=14, value=row.processed_on)
        self.ws.cell(row=r, column=15, value=row.status)
        self.ws.cell(row=r, column=16, value=row.bank_ref)
        row.sheet_row = r

    def append(self, rows: list[PurchaseRow]) -> None:
        """Append after the last data row; sets each row.sheet_row."""
        r = self._last_data_row + 1 if self._has_data() else FIRST_DATA_ROW
        for row in rows:
            self._write_row(r, row)
            r += 1
        self._last_data_row = r - 1

    def update(self, rows: list[PurchaseRow]) -> None:
        """Rewrite rows in place by sheet_row (used by reconcile for bank_ref/notes/status)."""
        for row in rows:
            if row.sheet_row is None:
                raise ValueError("update() requires row.sheet_row to be set")
            self._write_row(row.sheet_row, row)

    # ------------------------------------------------------------------
    # Bank sheet
    # ------------------------------------------------------------------

    def _find_last_bank_row(self) -> int:
        row = BANK_FIRST_DATA_ROW
        last = BANK_HEADER_ROW
        while True:
            val = self.ws_bank.cell(row=row, column=BANK_FIRST_COL).value
            if val is None:
                break
            last = row
            row += 1
        return last

    def _has_bank_data(self) -> bool:
        return self._last_bank_row is not None and self._last_bank_row >= BANK_FIRST_DATA_ROW

    def _ensure_bank_sheet(self) -> None:
        if self.ws_bank is not None:
            return
        self.ws_bank = self.wb.create_sheet(BankRow.SHEET_NAME)
        self.ws_bank.cell(row=1, column=BANK_FIRST_COL, value=BANK_TITLE)
        for i, header in enumerate(BankRow.COLUMNS):
            self.ws_bank.cell(row=BANK_HEADER_ROW, column=BANK_FIRST_COL + i, value=header)
        self._last_bank_row = BANK_HEADER_ROW

    def bank_rows(self) -> list[BankRow]:
        """All Bank-sheet data rows, with sheet_row set. Empty list if the sheet doesn't exist yet."""
        result: list[BankRow] = []
        if self.ws_bank is None or not self._has_bank_data():
            return result
        for r in range(BANK_FIRST_DATA_ROW, self._last_bank_row + 1):
            ws = self.ws_bank
            row_date = ws.cell(row=r, column=2).value
            if isinstance(row_date, datetime):
                row_date = row_date.date()
            description = ws.cell(row=r, column=3).value or ""
            amount = ws.cell(row=r, column=4).value
            direction = ws.cell(row=r, column=5).value or "debit"
            status = ws.cell(row=r, column=6).value or "No receipt"
            matched_source = ws.cell(row=r, column=7).value or ""
            ref = ws.cell(row=r, column=8).value or ""
            statement_file = ws.cell(row=r, column=9).value or ""
            statement_start = ws.cell(row=r, column=10).value
            if isinstance(statement_start, datetime):
                statement_start = statement_start.date()
            statement_end = ws.cell(row=r, column=11).value
            if isinstance(statement_end, datetime):
                statement_end = statement_end.date()
            processed_on = ws.cell(row=r, column=12).value

            result.append(
                BankRow(
                    date=row_date,
                    description=str(description),
                    amount=Decimal(str(amount)) if amount is not None else Decimal("0"),
                    direction=direction,
                    status=status,
                    matched_source=str(matched_source),
                    ref=str(ref),
                    statement_file=str(statement_file),
                    statement_start=statement_start,
                    statement_end=statement_end,
                    processed_on=processed_on,
                    sheet_row=r,
                )
            )
        return result

    def _write_bank_row(self, r: int, row: BankRow) -> None:
        ws = self.ws_bank
        ws.cell(row=r, column=2, value=row.date)
        ws.cell(row=r, column=3, value=row.description)
        ws.cell(row=r, column=4, value=float(round(Decimal(row.amount), 2)))
        ws.cell(row=r, column=5, value=row.direction)
        ws.cell(row=r, column=6, value=row.status)
        ws.cell(row=r, column=7, value=row.matched_source)
        ws.cell(row=r, column=8, value=row.ref)
        ws.cell(row=r, column=9, value=row.statement_file)
        ws.cell(row=r, column=10, value=row.statement_start)
        ws.cell(row=r, column=11, value=row.statement_end)
        ws.cell(row=r, column=12, value=row.processed_on)
        row.sheet_row = r

    def append_bank(self, rows: list[BankRow]) -> None:
        """Append after the last Bank-sheet data row (creating the sheet if needed); sets each
        row.sheet_row."""
        self._ensure_bank_sheet()
        r = self._last_bank_row + 1 if self._has_bank_data() else BANK_FIRST_DATA_ROW
        for row in rows:
            self._write_bank_row(r, row)
            r += 1
        self._last_bank_row = r - 1

    def update_bank(self, rows: list[BankRow]) -> None:
        """Rewrite Bank-sheet rows in place by sheet_row."""
        self._ensure_bank_sheet()
        for row in rows:
            if row.sheet_row is None:
                raise ValueError("update_bank() requires row.sheet_row to be set")
            self._write_bank_row(row.sheet_row, row)

    # ------------------------------------------------------------------
    # Save
    # ------------------------------------------------------------------

    def save(self) -> None:
        """Re-apply formats, fills, Check formulas, table range and totals row over ALL rows (all
        sheets), then write atomically."""
        self._save_purchases_sheet()
        if self.ws_bank is not None:
            self._save_bank_sheet()

        # Atomic write
        self.path.parent.mkdir(parents=True, exist_ok=True)
        tmp_path = self.path.with_suffix(".xlsx.tmp")
        self.wb.save(tmp_path)
        os.replace(tmp_path, self.path)

    def _save_purchases_sheet(self) -> None:
        ws = self.ws

        # Header row bold
        for i in range(len(PurchaseRow.COLUMNS)):
            ws.cell(row=HEADER_ROW, column=FIRST_COL + i).font = Font(bold=True)

        has_data = self._has_data()
        last = self._last_data_row

        if has_data:
            for r in range(FIRST_DATA_ROW, last + 1):
                ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
                ws.cell(row=r, column=5).number_format = "#,##0.00"
                ws.cell(row=r, column=6).number_format = "#,##0.00"
                ws.cell(row=r, column=7).number_format = "#,##0.00"
                ws.cell(row=r, column=8).number_format = "0.000"
                ws.cell(row=r, column=14).number_format = "yyyy-mm-dd hh:mm"
                ws.cell(row=r, column=10, value=_check_formula(r))

                status = ws.cell(row=r, column=15).value
                if status == "Review":
                    fill = REVIEW_FILL
                    font = NORMAL_FONT
                elif status == "Superseded":
                    fill = SUPERSEDED_FILL
                    font = SUPERSEDED_FONT
                else:
                    fill = NO_FILL
                    font = NORMAL_FONT
                for c in range(FIRST_COL, LAST_COL + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.fill = fill
                    cell.font = font

        # Column widths
        for col_letter, width in COLUMN_WIDTHS.items():
            ws.column_dimensions[col_letter].width = width

        # Table
        table_name = PurchaseRow.TABLE_NAME
        if table_name in ws.tables:
            del ws.tables[table_name]

        if has_data:
            table_ref = f"B{HEADER_ROW}:P{last}"
        else:
            # openpyxl needs at least one data row; use a blank row, skip totals.
            table_ref = f"B{HEADER_ROW}:P{FIRST_DATA_ROW}"

        table = Table(displayName=table_name, ref=table_ref)
        table.tableStyleInfo = _TABLE_STYLE
        ws.add_table(table)

        # Totals row (skip when zero data rows)
        if has_data:
            totals_row = last + 1
            # Clear any stray content below (in case previous totals row was longer/shorter)
            ws.cell(row=totals_row, column=2, value="Total")
            for c in range(3, 5):
                ws.cell(row=totals_row, column=c, value=None)
            ws.cell(row=totals_row, column=5, value=f"=SUBTOTAL(109,E{FIRST_DATA_ROW}:E{last})")
            ws.cell(row=totals_row, column=6, value=f"=SUBTOTAL(109,F{FIRST_DATA_ROW}:F{last})")
            ws.cell(row=totals_row, column=7, value=f"=SUBTOTAL(109,G{FIRST_DATA_ROW}:G{last})")
            ws.cell(row=totals_row, column=5).number_format = "#,##0.00"
            ws.cell(row=totals_row, column=6).number_format = "#,##0.00"
            ws.cell(row=totals_row, column=7).number_format = "#,##0.00"
            for c in range(8, LAST_COL + 1):
                ws.cell(row=totals_row, column=c, value=None)

        ws.freeze_panes = "B4"

    def _save_bank_sheet(self) -> None:
        ws = self.ws_bank

        for i in range(len(BankRow.COLUMNS)):
            ws.cell(row=BANK_HEADER_ROW, column=BANK_FIRST_COL + i).font = Font(bold=True)

        has_data = self._has_bank_data()
        last = self._last_bank_row

        if has_data:
            for r in range(BANK_FIRST_DATA_ROW, last + 1):
                ws.cell(row=r, column=2).number_format = "yyyy-mm-dd"
                ws.cell(row=r, column=4).number_format = "#,##0.00"
                ws.cell(row=r, column=10).number_format = "yyyy-mm-dd"
                ws.cell(row=r, column=11).number_format = "yyyy-mm-dd"
                ws.cell(row=r, column=12).number_format = "yyyy-mm-dd hh:mm"

                status = ws.cell(row=r, column=6).value
                if status == "No receipt":
                    fill = REVIEW_FILL
                elif status == "Ignored":
                    fill = SUPERSEDED_FILL
                else:
                    fill = NO_FILL
                for c in range(BANK_FIRST_COL, BANK_LAST_COL + 1):
                    ws.cell(row=r, column=c).fill = fill

        table_name = BankRow.TABLE_NAME
        if table_name in ws.tables:
            del ws.tables[table_name]

        if has_data:
            table_ref = f"B{BANK_HEADER_ROW}:L{last}"
        else:
            table_ref = f"B{BANK_HEADER_ROW}:L{BANK_FIRST_DATA_ROW}"

        table = Table(displayName=table_name, ref=table_ref)
        table.tableStyleInfo = _TABLE_STYLE
        ws.add_table(table)

        ws.freeze_panes = "B4"

